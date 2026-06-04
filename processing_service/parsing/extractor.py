import io
import logging
import os
from dataclasses import dataclass, field

logger = logging.getLogger(__name__)

_MIN_GOOD_CHARS = 100  # fewer extracted chars → low_confidence or failed

# Optional heavy imports — absent means that path degrades gracefully
try:
    import pdfplumber
    _HAS_PDFPLUMBER = True
except ImportError:  # pragma: no cover
    pdfplumber = None  # type: ignore[assignment]
    _HAS_PDFPLUMBER = False

try:
    import pytesseract
    from pdf2image import convert_from_bytes as _pdf2images
    _HAS_OCR = True
except ImportError:
    pytesseract = None  # type: ignore[assignment]
    _pdf2images = None  # type: ignore[assignment]
    _HAS_OCR = False

try:
    from docx import Document as _DocxDocument
    _HAS_DOCX = True
except ImportError:  # pragma: no cover
    _DocxDocument = None  # type: ignore[assignment]
    _HAS_DOCX = False


@dataclass
class ParseResult:
    text: str
    quality: str  # "good" | "low_confidence" | "failed"
    metadata: dict = field(default_factory=dict)


def extract(file_bytes: bytes, filename: str) -> ParseResult:
    """
    Extract text from a PDF or DOCX file.
    Never raises — failures are captured in quality="failed".
    """
    ext = os.path.splitext(filename.lower())[1]

    if ext == ".pdf":
        return _extract_pdf(file_bytes, filename)
    if ext == ".docx":
        return _extract_docx(file_bytes, filename)

    logger.warning("unsupported file type: %s", filename)
    return ParseResult(
        text="",
        quality="failed",
        metadata={"reason": f"unsupported_type:{ext}"},
    )


# ── PDF ───────────────────────────────────────────────────────────────────────


def _extract_pdf(file_bytes: bytes, filename: str) -> ParseResult:
    page_count = 0

    if not _HAS_PDFPLUMBER:  # pragma: no cover
        return _ocr_pdf(file_bytes, filename, page_count=0)

    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            page_count = len(pdf.pages)
            pages_text = [page.extract_text() or "" for page in pdf.pages]

        text = "\n\n".join(p for p in pages_text if p.strip())

        if len(text.strip()) >= _MIN_GOOD_CHARS:
            return ParseResult(
                text=text,
                quality="good",
                metadata={"method": "pdfplumber", "pages": page_count, "char_count": len(text)},
            )

        # Text too short — likely scanned; attempt OCR
        logger.info("%s: pdfplumber extracted %d chars, falling back to OCR", filename, len(text.strip()))
        return _ocr_pdf(file_bytes, filename, page_count)

    except Exception as exc:
        logger.warning("pdfplumber failed on %s: %s", filename, exc)
        return _ocr_pdf(file_bytes, filename, page_count)


def _ocr_pdf(file_bytes: bytes, filename: str, page_count: int) -> ParseResult:
    if not _HAS_OCR:
        logger.warning("OCR unavailable (pytesseract/pdf2image not installed) for %s", filename)
        quality = "low_confidence" if page_count > 0 else "failed"
        return ParseResult(
            text="",
            quality=quality,
            metadata={"reason": "ocr_unavailable", "pages": page_count},
        )

    try:
        images = _pdf2images(file_bytes, dpi=200)
        pages_text = [pytesseract.image_to_string(img) for img in images]
        text = "\n\n".join(p for p in pages_text if p.strip())
        quality = "good" if len(text.strip()) >= _MIN_GOOD_CHARS else "low_confidence"
        return ParseResult(
            text=text,
            quality=quality,
            metadata={"method": "ocr", "pages": len(images), "char_count": len(text)},
        )
    except Exception as exc:
        logger.warning("OCR failed on %s: %s", filename, exc)
        return ParseResult(
            text="",
            quality="failed",
            metadata={"reason": f"ocr_error:{exc}"},
        )


# ── DOCX ──────────────────────────────────────────────────────────────────────


def _extract_docx(file_bytes: bytes, filename: str) -> ParseResult:
    if not _HAS_DOCX:  # pragma: no cover
        return ParseResult(text="", quality="failed", metadata={"reason": "python-docx not installed"})

    try:
        doc = _DocxDocument(io.BytesIO(file_bytes))

        parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]

        # Also pull text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text.strip())

        text = "\n\n".join(parts)
        quality = "good" if len(text.strip()) >= _MIN_GOOD_CHARS else "low_confidence"
        return ParseResult(
            text=text,
            quality=quality,
            metadata={"method": "python-docx", "char_count": len(text)},
        )
    except Exception as exc:
        logger.warning("DOCX extraction failed on %s: %s", filename, exc)
        return ParseResult(
            text="",
            quality="failed",
            metadata={"reason": f"docx_error:{exc}"},
        )
